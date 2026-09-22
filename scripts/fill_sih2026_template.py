"""Fill official SIH 2026 idea template with MailTrace content.

Language / bullet style follows selected SIH decks (e.g. DevWise):
short points, plain sentences, Social / Economic / Environmental impact.

Keeps official branding and required section pointers. Expands the body
text boxes and drops font size so content fits above the footer (winning
teams do this; the stock boxes are too small for real idea text).

Slide 7 (instructions) is dropped for portal upload.
"""

from copy import deepcopy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SRC = Path(r"C:\Users\prajj\Downloads\SIH2026-IDEA-Presentation-Format (2).pptx")
DST = Path(r"C:\Users\prajj\Documents\proto.SIH\MailTrace_SIH26106_SIH2026.pptx")
TEAM = "prajjwaldubey"

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
INK = RGBColor(0x1A, 0x1A, 0x1A)
HEAD = RGBColor(0x0B, 0x2F, 0x6B)


def set_runs_text(paragraph, text: str, size_pt: float | None = None, bold: bool | None = None, color=None) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
        runs = [paragraph.runs[0]]
    else:
        paragraph.text = text
        runs = list(paragraph.runs) or []
    for run in runs:
        if size_pt is not None:
            run.font.size = Pt(size_pt)
        if bold is not None:
            run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color
        run.font.name = "Calibri"


def clone_paragraph(text_frame, source_para):
    new_p = deepcopy(source_para._p)
    text_frame._txBody.append(new_p)
    return text_frame.paragraphs[-1]


def clear_bullet(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("buFont", "buChar", "buAutoNum", "buBlip", "buClr", "buSzPct", "buSzPts"):
        el = pPr.find(qn(f"a:{tag}"))
        if el is not None:
            pPr.remove(el)
    if pPr.find(qn("a:buNone")) is None:
        etree.SubElement(pPr, f"{{{A_NS}}}buNone")


def ensure_bullet(paragraph, char: str = "•") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("buNone", "buAutoNum", "buBlip"):
        el = pPr.find(qn(f"a:{tag}"))
        if el is not None:
            pPr.remove(el)
    bu = pPr.find(qn("a:buChar"))
    if bu is None:
        bu = etree.SubElement(pPr, f"{{{A_NS}}}buChar")
    bu.set("char", char)


def set_spacing(paragraph, before=0, after=4) -> None:
    """Set paragraph spacing via OOXML (works on cloned paras)."""
    pPr = paragraph._p.get_or_add_pPr()
    spc_bef = pPr.find(qn("a:spcBef"))
    if spc_bef is not None:
        pPr.remove(spc_bef)
    spc_aft = pPr.find(qn("a:spcAft"))
    if spc_aft is not None:
        pPr.remove(spc_aft)
    if before:
        el = etree.SubElement(pPr, f"{{{A_NS}}}spcBef")
        etree.SubElement(el, f"{{{A_NS}}}spcPts").set("val", str(int(before * 100)))
    if after:
        el = etree.SubElement(pPr, f"{{{A_NS}}}spcAft")
        etree.SubElement(el, f"{{{A_NS}}}spcPts").set("val", str(int(after * 100)))


def expand_box(shape, top_in: float, height_in: float, left_in: float | None = None, width_in: float | None = None) -> None:
    shape.top = Inches(top_in)
    shape.height = Inches(height_in)
    if left_in is not None:
        shape.left = Inches(left_in)
    if width_in is not None:
        shape.width = Inches(width_in)


def fill_box(shape, lines: list[tuple[str, str]], body_pt: float = 13, head_pt: float = 15) -> None:
    """lines: (kind, text) where kind is 'h' (heading) or 'b' (bullet)."""
    tf = shape.text_frame
    tf.word_wrap = True
    paras = list(tf.paragraphs)
    if not paras:
        return
    style = paras[min(len(paras) - 1, 3)]
    for i, (kind, text) in enumerate(lines):
        if i < len(paras):
            para = paras[i]
        else:
            para = clone_paragraph(tf, style)
        if kind == "h":
            set_runs_text(para, text, size_pt=head_pt, bold=True, color=HEAD)
            clear_bullet(para)
            set_spacing(para, before=6 if i else 0, after=2)
        else:
            set_runs_text(para, text, size_pt=body_pt, bold=False, color=INK)
            ensure_bullet(para, "•")
            set_spacing(para, before=0, after=3)
    for j in range(len(lines), len(list(tf.paragraphs))):
        if j < len(paras):
            set_runs_text(paras[j], "")
            clear_bullet(paras[j])


def shape_by_name(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def oval_team(slide):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.text_frame.text.strip() in ("Your Team Name", TEAM):
            # keep oval readable for a long portal name
            set_runs_text(sh.text_frame.paragraphs[0], TEAM, size_pt=9, bold=True)
            try:
                sh.width = Inches(1.55)
                sh.height = Inches(0.55)
            except Exception:
                pass


def delete_slide(prs: Presentation, index: int) -> None:
    rId = prs.slides._sldIdLst[index].get(qn("r:id"))
    prs.part.drop_rel(rId)
    sldId = prs.slides._sldIdLst[index]
    prs.slides._sldIdLst.remove(sldId)


def main() -> None:
    prs = Presentation(str(SRC))
    s1, s2, s3, s4, s5, s6 = (prs.slides[i] for i in range(6))

    # --- Slide 1 title page ---
    tb1 = shape_by_name(s1, "TextBox 9")
    title_lines = [
        "",
        "Problem Statement ID – SIH26106",
        "Problem Statement Title – AI-Powered Email Threat Detection, GeoLocation and Forensic Intelligence Platform",
        "Theme – Blockchain & Cybersecurity",
        "PS Category – Software",
        "Team ID – [Fill from SIH portal]",
        "Team Name (Registered on portal) – prajjwaldubey",
    ]
    tf = tb1.text_frame
    tf.word_wrap = True
    paras = list(tf.paragraphs)
    for i, line in enumerate(title_lines):
        if i < len(paras):
            para = paras[i]
        else:
            para = clone_paragraph(tf, paras[-1])
        # long PS title needs slightly smaller type so Category / Team stay on slide
        size = 15 if i == 2 else 16
        set_runs_text(para, line, size_pt=size if line else 12, bold=bool(line), color=INK)
        set_spacing(para, before=4, after=8)
    for j in range(len(title_lines), len(list(tf.paragraphs))):
        if j < len(paras):
            set_runs_text(paras[j], "")

    # --- Slide 2 idea ---
    t2 = shape_by_name(s2, "Title 1")
    set_runs_text(t2.text_frame.paragraphs[0], "MailTrace", size_pt=32, bold=True)
    box2 = shape_by_name(s2, "TextBox 8")
    expand_box(box2, top_in=1.55, height_in=5.35, left_in=0.45, width_in=12.4)
    fill_box(
        box2,
        [
            ("h", "Proposed Solution:"),
            (
                "b",
                "Develop a comprehensive web platform that analyses original emails (paste or .eml upload) to detect threats, map mail-server locations, and prepare a forensic report for institutes and SOC teams.",
            ),
            (
                "b",
                "Features include SPF / DKIM / DMARC checks, scam-language cues, lookalike domain spotting, hop-path geolocation, risk score with reasons, and a downloadable PDF case report.",
            ),
            ("h", "How it Addresses the Problem:"),
            (
                "b",
                "Detection Engine: Spots phishing, BEC, and spoofed mail using headers + body patterns, helping institutes catch fraud before users click.",
            ),
            (
                "b",
                "GeoLocation Module: Shows where mail servers sat along the path — useful for investigation without naming a person’s home address.",
            ),
            (
                "b",
                "Forensic Report: Gives analysts a ready PDF with score, reasons, stamps, and map for tickets and incident response.",
            ),
            ("h", "Innovation and Uniqueness:"),
            ("b", "End-to-end flow: detection + map + PDF in one laptop tool."),
            ("b", "Explainable score — every point has a written reason."),
            ("b", "Honest geo: mail servers / ISP only, not personal identity."),
            ("b", "Works on real Gmail “Show original” and .eml files; no mailbox login."),
            ("b", "Graph of domains, IPs, and clues for campaign-style review."),
        ],
        body_pt=13,
        head_pt=15,
    )
    oval_team(s2)

    # --- Slide 3 technical ---
    box3 = shape_by_name(s3, "TextBox 8")
    expand_box(box3, top_in=1.7, height_in=5.15, left_in=0.55, width_in=12.2)
    fill_box(
        box3,
        [
            ("h", "Tech-Stack"),
            ("b", "Frontend: HTML, CSS, JavaScript with Leaflet map for hop visualisation."),
            ("b", "Backend: Python 3, FastAPI, Uvicorn — runs as a local / cloud web service."),
            (
                "b",
                "Analysis: RFC822 email parser, rule + NLP cue engine, SPF / DKIM / DMARC from Authentication-Results.",
            ),
            (
                "b",
                "Geo & storage: Demo IP table + optional public IP lookup; SQLite case store; ReportLab forensic PDF.",
            ),
            ("h", "Methodology and Process:"),
            ("b", "Ingest raw email → parse headers, body, URLs, and attachments."),
            ("b", "Validate sender stamps → extract public IPv4 from Received: lines → geolocate hops."),
            ("b", "Score content + identity mismatches → show banner, reasons, and map → export PDF."),
            ("b", "Working prototype already runs on a laptop; same pipeline for genuine and malicious sample mails."),
        ],
        body_pt=14,
        head_pt=16,
    )
    oval_team(s3)

    # --- Slide 4 feasibility ---
    box4 = shape_by_name(s4, "TextBox 8")
    expand_box(box4, top_in=1.7, height_in=5.15, left_in=0.55, width_in=12.2)
    fill_box(
        box4,
        [
            ("h", "Feasibility"),
            ("b", "Software-only problem: no RF hardware, no classified dataset, no Gmail password needed."),
            ("b", "Feasible as a student laptop demo; real “Show original” mails already pass through the same pipeline."),
            ("h", "Challenges and Risks"),
            ("b", "Gmail may hide origin IP / show IPv6-only hops → map can look empty on Gmail–Gmail mail."),
            ("b", "No official SIH phishing dataset; CDN or font links on real newsletters can add small score noise."),
            ("b", "Risk of over-claiming a deep-learning model when the build uses explainable rules + cues."),
            ("h", "Overcoming Challenges"),
            ("b", "Empty maps are captioned as expected when no public IPv4 is present; score and stamps still run."),
            ("b", "Demo uses documented lab traces plus real pasted originals; reason list stays open for human review."),
            ("b", "Position AI as explainable detection now; same features can feed a trained classifier later."),
            ("b", "Ethics: analyse only user-provided mail; attribute infrastructure, not a named person."),
        ],
        body_pt=13.5,
        head_pt=15,
    )
    oval_team(s4)

    # --- Slide 5 impact ---
    box5 = shape_by_name(s5, "TextBox 8")
    expand_box(box5, top_in=1.7, height_in=5.15, left_in=0.55, width_in=12.2)
    fill_box(
        box5,
        [
            ("h", "Potential Impact on the Target Audience:"),
            (
                "b",
                "Improved threat awareness and faster triage of fake dean, GST, invoice, and internship-scam mail for college / institute SOC and IT admin.",
            ),
            (
                "b",
                "Empowerment of students and staff with a clear “do not pay / click” vs “looks legitimate” view, and reasons they can follow.",
            ),
            (
                "b",
                "Enhanced readiness for CERT-style review of a single suspicious message without reading raw headers by hand.",
            ),
            ("h", "Benefits of the Solution:"),
            (
                "b",
                "Social: Bridges the gap between spam filters and real investigation; promotes safer campus email use.",
            ),
            (
                "b",
                "Economic: Prepares teams to cut fraud loss and analyst time per ticket; PDF supports internal records.",
            ),
            ("b", "Environmental: Reduces paper usage with digital forensic reports instead of printed header dumps."),
        ],
        body_pt=14,
        head_pt=16,
    )
    oval_team(s5)

    # --- Slide 6 research ---
    box6 = shape_by_name(s6, "TextBox 8")
    expand_box(box6, top_in=1.7, height_in=5.15, left_in=0.55, width_in=12.2)
    fill_box(
        box6,
        [
            ("h", "Email Authentication Standards"),
            ("b", "RFC 7208 — SPF: how receivers check if a host may send for a domain. https://datatracker.ietf.org/doc/html/rfc7208"),
            ("b", "RFC 6376 — DKIM: cryptographic signature so body/headers can be verified. https://datatracker.ietf.org/doc/html/rfc6376"),
            ("b", "RFC 7489 — DMARC: policy aligning SPF/DKIM with the visible From domain. https://datatracker.ietf.org/doc/html/rfc7489"),
            ("h", "Threat Landscape and Practice"),
            ("b", "CERT-In / MeitY advisories: guidance on phishing and business email compromise in India. https://www.cert-in.org.in/"),
            ("b", "APWG Phishing Activity Trends: shows volume of credential and payment fraud mail. https://apwg.org/"),
            ("b", "RFC 5321 — SMTP: mail is relayed; Received chain is servers, not a person’s house. https://datatracker.ietf.org/doc/html/rfc5321"),
        ],
        body_pt=13,
        head_pt=15,
    )
    oval_team(s6)

    if len(prs.slides) >= 7:
        delete_slide(prs, 6)

    prs.save(str(DST))
    print(DST)


if __name__ == "__main__":
    main()

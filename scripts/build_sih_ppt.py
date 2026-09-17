"""Generate SIH26106 idea-submission PPT (DevWise 6-slide layout)."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
TEAL = RGBColor(0x1A, 0xB3, 0xA6)
GOLD = RGBColor(0xE8, 0xA8, 0x17)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x2B, 0x3A)
MUTED = RGBColor(0x5A, 0x6B, 0x7C)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
CARD = RGBColor(0xE8, 0xEE, 0xF4)

W, H = Inches(13.333), Inches(7.5)


def set_run(run, size, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def add_text(slide, l, t, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return box


def bullets(slide, l, t, w, h, items, size=14, color=INK, spacing=6):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(spacing)
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
    return box


def header_bar(slide, kicker, title):
    add_rect(slide, 0, 0, W, Inches(1.05), NAVY)
    add_rect(slide, 0, Inches(1.05), W, Inches(0.08), TEAL)
    add_text(slide, Inches(0.4), Inches(0.12), Inches(12.4), Inches(0.32), kicker, 11, False, TEAL)
    add_text(slide, Inches(0.4), Inches(0.38), Inches(12.4), Inches(0.55), title, 26, True, WHITE)


def footer(slide, n):
    add_rect(slide, 0, Inches(7.22), W, Inches(0.28), NAVY)
    add_text(slide, Inches(0.4), Inches(7.24), Inches(10), Inches(0.24), "SIH Idea submission  ·  SIH26106 MailTrace  ·  Software prototype", 10, False, WHITE)
    add_text(slide, Inches(11.4), Inches(7.24), Inches(1.5), Inches(0.24), f"{n} / 6", 10, False, TEAL, PP_ALIGN.RIGHT)


def slide1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, NAVY)
    add_rect(s, 0, 0, Inches(0.18), H, TEAL)
    add_text(s, Inches(0.55), Inches(0.35), Inches(12), Inches(0.35), "SMART INDIA HACKATHON", 14, True, TEAL)
    add_text(s, Inches(0.55), Inches(0.75), Inches(12.2), Inches(1.15), "MailTrace", 48, True, WHITE)
    add_text(
        s,
        Inches(0.55),
        Inches(1.85),
        Inches(12.2),
        Inches(0.7),
        "AI-Powered Email Threat Detection, GeoLocation\nand Forensic Intelligence Platform",
        20,
        False,
        GOLD,
    )

    rows = [
        ("Problem Statement ID", "SIH26106  (confirm on official portal)"),
        ("Problem Statement Title", "AI-Powered Email Threat Detection, GeoLocation and Forensic Intelligence Platform"),
        ("Organisation", "AICTE Cyber Security Cell"),
        ("PS Category", "Software"),
        ("Theme", "[Fill from SIH portal]"),
        ("Team ID", "[Fill from SIH portal]"),
        ("Team Name", "prajjwaldubey"),
    ]
    y = Inches(2.7)
    for label, val in rows:
        add_rect(s, Inches(0.55), y, Inches(12.2), Inches(0.52), RGBColor(0x12, 0x2B, 0x4A))
        add_text(s, Inches(0.7), y + Inches(0.04), Inches(3.3), Inches(0.42), label, 12, True, TEAL)
        add_text(s, Inches(4.1), y + Inches(0.04), Inches(8.4), Inches(0.42), val, 13, False, WHITE)
        y += Inches(0.56)
    add_text(s, Inches(0.55), Inches(6.75), Inches(12), Inches(0.35), "Prototype  ·  Analyse original email  ·  Locate mail servers, not a person", 13, False, TEAL)


def slide2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, LIGHT)
    header_bar(s, "SIH IDEA SUBMISSION", "Proposed Solution")
    # three cards
    add_rect(s, Inches(0.35), Inches(1.35), Inches(4.1), Inches(5.65), WHITE)
    add_rect(s, Inches(0.35), Inches(1.35), Inches(4.1), Inches(0.42), TEAL)
    add_text(s, Inches(0.5), Inches(1.4), Inches(3.8), Inches(0.35), "Proposed solution", 14, True, WHITE)
    bullets(
        s,
        Inches(0.5),
        Inches(1.9),
        Inches(3.8),
        Inches(4.9),
        [
            "Ingest: paste Show original or upload .eml",
            "Threat score 0–100 from SPF/DKIM/DMARC + scam cues (urgency, lookalike, bad URLs)",
            "Labels: legitimate / suspicious / phishing-BEC",
            "Map: public IPv4 hops from Received: headers",
            "Forensic PDF + saved case list for SOC/admin",
        ],
        13,
        spacing=10,
    )

    add_rect(s, Inches(4.6), Inches(1.35), Inches(4.1), Inches(5.65), WHITE)
    add_rect(s, Inches(4.6), Inches(1.35), Inches(4.1), Inches(0.42), NAVY)
    add_text(s, Inches(4.75), Inches(1.4), Inches(3.8), Inches(0.35), "How it addresses the PS", 14, True, WHITE)
    bullets(
        s,
        Inches(4.75),
        Inches(1.9),
        Inches(3.8),
        Inches(4.9),
        [
            "PS asks for detection + geolocation + investigative report — one laptop app",
            "Same pipeline for genuine and malicious mail",
            "Every point is a logged reason (not a black box)",
            "User provides the message; no mailbox intrusion",
            "Gmail–Gmail empty map is explained, not hidden",
        ],
        13,
        spacing=10,
    )

    add_rect(s, Inches(8.85), Inches(1.35), Inches(4.1), Inches(5.65), WHITE)
    add_rect(s, Inches(8.85), Inches(1.35), Inches(4.1), Inches(0.42), GOLD)
    add_text(s, Inches(9.0), Inches(1.4), Inches(3.8), Inches(0.35), "Uniqueness", 14, True, NAVY)
    bullets(
        s,
        Inches(9.0),
        Inches(1.9),
        Inches(3.8),
        Inches(4.9),
        [
            "Explainable scoring for forensics",
            "Hop map + PDF in one flow",
            "Honest attribution: servers / ISP, not a named person",
            "Works on real Show original mails",
            "Heuristic engine now; same features ready for ML later",
        ],
        13,
        spacing=10,
    )
    footer(s, 2)


def slide3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, LIGHT)
    header_bar(s, "TECHNICAL APPROACH", "Architecture and tech stack")

    steps = [
        ("1. Ingest", "Raw RFC822 / .eml"),
        ("2. Parse", "Headers, body, URLs"),
        ("3. Auth", "SPF · DKIM · DMARC"),
        ("4. Heuristics", "Cues + weights"),
        ("5. Geo", "Public IPv4 lookup"),
        ("6. Output", "Score · map · PDF"),
    ]
    x = Inches(0.35)
    for title, sub in steps:
        add_rect(s, x, Inches(1.4), Inches(2.0), Inches(1.35), WHITE)
        add_rect(s, x, Inches(1.4), Inches(2.0), Inches(0.12), TEAL)
        add_text(s, x + Inches(0.1), Inches(1.6), Inches(1.8), Inches(0.45), title, 14, True, NAVY)
        add_text(s, x + Inches(0.1), Inches(2.1), Inches(1.8), Inches(0.5), sub, 12, False, MUTED)
        x += Inches(2.15)

    add_rect(s, Inches(0.35), Inches(3.0), Inches(6.3), Inches(3.95), WHITE)
    add_text(s, Inches(0.55), Inches(3.15), Inches(6), Inches(0.35), "Tech stack", 16, True, NAVY)
    bullets(
        s,
        Inches(0.55),
        Inches(3.55),
        Inches(5.9),
        Inches(3.2),
        [
            "Python 3 · FastAPI · Uvicorn (local prototype)",
            "Email parser (RFC822) · rule engine for threat score",
            "Geo: demo IP table + optional ip-api for unknown IPs",
            "SQLite case store · ReportLab forensic PDF",
            "Leaflet + OpenStreetMap hop visualisation",
            "No Gmail login, no classified dataset required",
        ],
        14,
        spacing=8,
    )

    add_rect(s, Inches(6.85), Inches(3.0), Inches(6.1), Inches(3.95), WHITE)
    add_text(s, Inches(7.05), Inches(3.15), Inches(5.7), Inches(0.35), "What we do not claim (by design)", 16, True, NAVY)
    bullets(
        s,
        Inches(7.05),
        Inches(3.55),
        Inches(5.7),
        Inches(3.2),
        [
            "Not a trained deep-learning model in this build",
            "Not GPS / home address of a human sender",
            "Not live DNS re-query of SPF (uses Authentication-Results on the mail)",
            "IPv6-only hops are not plotted (IPv4 public IPs only)",
            "Next step: classifier on the same feature vector",
        ],
        14,
        spacing=8,
    )
    footer(s, 3)


def slide4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, LIGHT)
    header_bar(s, "FEASIBILITY AND VIABILITY", "Challenges, risks, and how we overcome them")

    headers = ["Challenge / risk", "Feasibility", "How we overcome it"]
    xs = [Inches(0.35), Inches(4.55), Inches(8.75)]
    for x, htxt in zip(xs, headers):
        add_rect(s, x, Inches(1.35), Inches(4.05), Inches(0.45), NAVY)
        add_text(s, x + Inches(0.12), Inches(1.4), Inches(3.8), Inches(0.35), htxt, 14, True, WHITE)

    rows = [
        (
            "Gmail hides sender IP; IPv6-only Received lines",
            "Expected for Gmail–Gmail; still a valid mail",
            "Show empty map with an explicit caption; score/stamps still run",
        ),
        (
            "No official SIH phishing dataset",
            "PS is software on user-provided mail",
            "Demo with real Show original + documented lab traces",
        ),
        (
            "Judges ask “where is the AI?”",
            "Heuristic engine is shippable now",
            "Explainable weighted features; ML can use the same inputs later",
        ),
        (
            "CDN/font links inflate score on real newsletters",
            "False nits are visible",
            "Reason list is open; human remains in the loop",
        ),
        (
            "Legal / ethics of tracing",
            "High — analyse only pasted mail",
            "Infrastructure attribution only; no identity claim",
        ),
    ]
    y = Inches(1.9)
    for i, (a, b, c) in enumerate(rows):
        bg = WHITE if i % 2 == 0 else CARD
        for x, txt in zip(xs, (a, b, c)):
            add_rect(s, x, y, Inches(4.05), Inches(0.95), bg)
            add_text(s, x + Inches(0.1), y + Inches(0.08), Inches(3.85), Inches(0.82), txt, 12, False, INK)
        y += Inches(0.98)
    footer(s, 4)


def slide5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, LIGHT)
    header_bar(s, "IMPACT AND BENEFITS", "Who this helps and why it matters")

    add_rect(s, Inches(0.35), Inches(1.4), Inches(12.6), Inches(1.35), WHITE)
    add_text(s, Inches(0.55), Inches(1.5), Inches(12.2), Inches(0.3), "Target audience", 16, True, NAVY)
    add_text(
        s,
        Inches(0.55),
        Inches(1.9),
        Inches(12.2),
        Inches(0.7),
        "College / institute SOC and IT admin · students and staff receiving fake dean, GST, or invoice mail · CERT-style review of a single suspicious message.",
        15,
        False,
        INK,
    )

    cards = [
        ("Social", "Fewer people click fake invoices or gift-card BEC. Clear “do not pay” vs “looks legitimate” in seconds."),
        ("Operational", "One screen instead of reading raw headers. Case history + PDF for tickets and internal review."),
        ("Trust & safety", "Prototype labelled honestly. We geolocate mail servers, not a person. No credential theft, no mailbox access."),
        ("National relevance", "Matches AICTE Cyber Security Cell need: threat detection + geo + forensic artefact on a laptop, no classified data."),
    ]
    positions = [
        (Inches(0.35), Inches(2.95)),
        (Inches(6.85), Inches(2.95)),
        (Inches(0.35), Inches(4.95)),
        (Inches(6.85), Inches(4.95)),
    ]
    for (x, y), (title, body) in zip(positions, cards):
        add_rect(s, x, y, Inches(6.15), Inches(1.85), WHITE)
        add_rect(s, x, y, Inches(0.12), Inches(1.85), TEAL)
        add_text(s, x + Inches(0.3), y + Inches(0.15), Inches(5.7), Inches(0.4), title, 16, True, NAVY)
        add_text(s, x + Inches(0.3), y + Inches(0.6), Inches(5.7), Inches(1.1), body, 13, False, INK)
    footer(s, 5)


def slide6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(s, 0, 0, W, H, LIGHT)
    header_bar(s, "RESEARCH AND REFERENCES", "Standards, advisories, and why geo ≠ identity")

    refs = [
        ("Email authentication (core of our stamps)", [
            "RFC 7208 — Sender Policy Framework (SPF)",
            "RFC 6376 — DomainKeys Identified Mail (DKIM)",
            "RFC 7489 — DMARC",
        ]),
        ("Threat landscape", [
            "CERT-In / MeitY advisories on phishing and BEC",
            "APWG Phishing Activity Trends — volume of credential and payment fraud mails",
        ]),
        ("Geolocation limits (we state these in the demo)", [
            "RFC 5321 — mail is relayed; Received chain is servers, not a person",
            "Gmail / mailbox providers strip or omit origin IPv4 — empty map is expected",
        ]),
        ("Forensic practice", [
            "Analyse user-provided original (Show original / .eml) only",
            "Infrastructure attribution + explainable score as a decision-support artefact",
        ]),
    ]
    y = Inches(1.3)
    for title, items in refs:
        add_text(s, Inches(0.5), y, Inches(12.3), Inches(0.32), title, 15, True, TEAL)
        y += Inches(0.32)
        for it in items:
            add_text(s, Inches(0.7), y, Inches(12.1), Inches(0.28), "•  " + it, 13, False, INK)
            y += Inches(0.28)
        y += Inches(0.12)

    add_rect(s, Inches(0.4), Inches(6.35), Inches(12.5), Inches(0.72), NAVY)
    add_text(
        s,
        Inches(0.55),
        Inches(6.45),
        Inches(12.2),
        Inches(0.52),
        "Demo: paste Show original → banner + stamps + reasons + map + PDF.  Prototype.  Servers, not a person.",
        14,
        True,
        WHITE,
    )
    footer(s, 6)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide1(prs)
    slide2(prs)
    slide3(prs)
    slide4(prs)
    slide5(prs)
    slide6(prs)
    out = r"C:\Users\prajj\Documents\proto.SIH\MailTrace_SIH26106.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
